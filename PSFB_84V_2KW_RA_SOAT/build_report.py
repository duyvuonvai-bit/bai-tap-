from pathlib import Path
import re,csv,json,hashlib,zipfile,subprocess,os
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import fitz
P=Path(__file__).resolve().parent
d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=s.bottom_margin=Cm(1.7);s.left_margin=s.right_margin=Cm(1.8)
normal=d.styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(10)
normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',20),('Heading 1',16),('Heading 2',13),('Heading 3',11)]:
 st=d.styles[name];st.font.name='DejaVu Sans';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('173F6B');st.paragraph_format.keep_with_next=True
foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;foot.add_run('PSFB 84 V • Bản rà soát, chưa có Gerber | ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
def rich(p,t):
 for v in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)',t):
  if not v:continue
  rr=p.add_run(v[2:-2] if v.startswith('**') else v[1:-1] if v.startswith('`') else v)
  if v.startswith('**'):rr.bold=True
def md(path):
 rows=path.read_text().splitlines();i=0
 while i<len(rows):
  t=rows[i].strip();i+=1
  if not t:continue
  if t.startswith('|'):
   data=[t]
   while i<len(rows) and rows[i].strip().startswith('|'):data.append(rows[i].strip());i+=1
   cells=[[x.strip() for x in q.strip('|').split('|')] for q in data if not re.fullmatch(r'[|\s:\-]+',q)]
   tab=d.add_table(rows=0,cols=len(cells[0]));tab.style='Light Shading Accent 1'
   widths=None
   if cells[0][0]=='Khối / linh kiện':widths=[3.7,1.8,4.7,7.2]
   elif cells[0][0]=='Khối trong ảnh':widths=[2.8,7.3,7.3]
   if widths:
    tab.autofit=False
    for col,width in zip(tab.columns,widths):col.width=Cm(width)
   for j,row in enumerate(cells):
    new=tab.add_row();pr=new._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
    if j==0:pr.append(OxmlElement('w:tblHeader'))
    for ci,(cell,txt) in enumerate(zip(new.cells,row)):
     if widths:cell.width=Cm(widths[ci])
     rich(cell.paragraphs[0],txt)
     for run in cell.paragraphs[0].runs:run.font.size=Pt(8.5);run.bold=j==0
   d.add_paragraph();continue
  if t.startswith('#'):
   h=len(t)-len(t.lstrip('#'));d.add_heading(t[h:].strip(),min(h,3))
  elif t.startswith('- '):rich(d.add_paragraph(style='List Bullet'),t[2:])
  elif re.match(r'^\d+\. ',t):rich(d.add_paragraph(style='List Number'),re.sub(r'^\d+\. ','',t))
  else:rich(d.add_paragraph(),t)

d.add_heading('BỘ SẠC PSFB 2 kW\n20S Li-ion • 84 V',0)
d.add_paragraph('Rà soát sơ đồ, tính thông số và bảng linh kiện có công dụng — 16/10/2026')
d.add_paragraph('Dùng bo STM32G474RET6 đã mua. Đã xác nhận mục tiêu cuối cùng là sơ đồ nguyên lý và PCB đặt gia công; giai đoạn hiện tại chưa có PCB/Gerber để chế tạo. Không hàn theo ảnh nguyên lý gốc vì có các lỗi kết nối và bảo vệ được chỉ rõ trong báo cáo.')
d.add_heading('Cấu hình thiết kế',2)
for t in ['220 VAC/50 Hz; PFC danh định400 V; dải bus tính380–420 V.','20S Li-ion4,2 V/cell; mục tiêu50 Ah, BMS dòng sạc liên tục≥30 A. Đây là cấu hình chọn để thiết kế, chưa xác minh pin thật.','CC tối đa23,81 A; CV84 V; 2 kW chỉ đạt ở84 V với giới hạn dòng này.','Bo MCU: sản phẩm Shopee1573220776/56315356231 do người dùng cung cấp; trang bị403, chưa xác minh đầu cắm/nguồn/kích thước.']:
 rich(d.add_paragraph(style='List Bullet'),t)
d.add_heading('Các phần trong tài liệu',2)
for t in ['1. Tính toán định lượng và các thông số còn phải xác minh.','2. Bảng linh kiện, số lượng và công dụng.','3. Rà soát chi tiết các lỗi trong ảnh nguyên lý.','4. Sơ đồ sửa các kết nối công suất chính (trang cuối).']:
 d.add_paragraph(t)
d.add_heading('Điều kiện còn thiếu để xuất PCB',2)
d.add_paragraph('Cần ảnh/pinout và kích thước đúng bo MCU; mã linh kiện và datasheet/footprint được xác minh; thiết kế từ tính với lõi thực; sơ đồ PFC, nguồn phụ, nguồn driver và bảo vệ nhanh/chốt lỗi đầy đủ. Sau đó mới kiểm ERC/DRC, xuất Gerber/drill/BOM và bản vẽ lắp. Bản báo cáo không tự biến các mục này thành thiết kế hoàn tất.')
for name in ['TINH_TOAN_SO_BO.md','BANG_LINH_KIEN_VA_CONG_DUNG.md','AUDIT_IMAGE.md']:
 d.add_page_break();md(P/name)
docx=P/'BAO_CAO_PSFB_84V_2KW.docx';d.save(docx)
env=os.environ.copy();env['XDG_CACHE_HOME']='/tmp/psfb_font_cache'
r=subprocess.run(['soffice','-env:UserInstallation=file:///tmp/psfb_lo_profile','--headless','--convert-to','pdf','--outdir',str(P),str(docx)],env=env,capture_output=True,text=True)
print(r.stdout);r.check_returncode()
pdf=P/'BAO_CAO_PSFB_84V_2KW.pdf';combined=fitz.open(pdf);ntext=len(combined)
with fitz.open(P/'SO_DO_KET_NOI_DA_SUA.pdf') as fig:combined.insert_pdf(fig)
combined.set_metadata({'title':'Rà soát và tính toán bộ sạc PSFB 2 kW, 84 V, STM32G474RET6','author':'','subject':'Tính sơ bộ; chưa có PCB/Gerber chế tạo'})
temp=P/'_merged.pdf';combined.save(temp,garbage=4,deflate=True);combined.close();temp.replace(pdf)
# Export the same function table as a spreadsheet-readable CSV.
lines=(P/'BANG_LINH_KIEN_VA_CONG_DUNG.md').read_text().splitlines()
rows=[[v.strip() for v in ln.strip('|').split('|')] for ln in lines if ln.startswith('|') and not re.fullmatch(r'[|\s:\-]+',ln)]
with (P/'BANG_LINH_KIEN_VA_CONG_DUNG.csv').open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows(rows)
print('Report pages',ntext+1)
